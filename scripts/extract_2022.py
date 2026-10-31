"""One-time extract of Lewis County's 2022 general election results into data/archive/20221108/.

Sources (not committed; download them to run this again):
  2022Gen_Precinct_Results_GIS-Ready.xlsx  Secretary of State precinct results for every county:
      registered voters, ballots cast, and the state, congressional, legislative and judicial
      contests. Statewide and district totals for the "vs. statewide" comparison come from it.
  2022-wa-local-precinct-general.csv       MIT Election Data and Science Lab (github.com/MEDSL/
      2022-elections-official): the county offices and local measures the state file leaves out,
      and candidates' party preferences.

Precinct boundaries: the state's 2022 shapefile isn't used. 2022's precincts are the same as the
August 2026 primary's except Winlock, which was later split into Winlock #1-#3, so this writes
site/data/20221108.geojson from site/data/20260804.geojson with those three merged back into one.

Usage:
  pip install openpyxl shapely
  python scripts/extract_2022.py <xlsx> <medsl csv>
"""

import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from shapely.geometry import mapping, shape
from shapely.ops import unary_union

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "archive" / "20221108"
SITE = ROOT / "site" / "data"
MERGED = {"WINLOCK": ["WINLOCK #1", "WINLOCK #2", "WINLOCK #3"]}

# State file contests on Lewis ballots: race key -> display name (as the 2026 feed words them).
STATE_CONTESTS = {
    ("SW", "Advisory Vote No. 39"): "Advisory Vote No. 39",
    ("SW", "Advisory Vote No. 40"): "Advisory Vote No. 40",
    ("FED", "U.S. Senator"): "U.S. Senator",
    ("CON", "U.S. Representative"): "U.S. Representative - Congressional District {n}",
    ("SW", "Secretary of State"): "Secretary of State",
    ("LEG", "State Representative Pos. 1"): "State Representative Pos. 1 - Legislative District {n}",
    ("LEG", "State Representative Pos. 2"): "State Representative Pos. 2 - Legislative District {n}",
    ("JUD", "Justice Position #01"): "Justice Position #01 - Supreme Court",
    ("JUD", "Justice Position #05"): "Justice Position #05 - Supreme Court",
    ("JUD", "Justice Position #06"): "Justice Position #06 - Supreme Court",
    ("JUD", "Judge Position 2"): "Judge Position 2 - Court of Appeals, Division 2, District 3",
}

# MEDSL offices the state file doesn't have -> display name.
LOCAL_CONTESTS = {
    "LEWIS ASSESSOR": "Assessor",
    "LEWIS AUDITOR": "Auditor",
    "LEWIS CLERK": "Clerk",
    "LEWIS CORONER": "Coroner",
    "LEWIS PROSECUTING ATTORNEY": "Prosecuting Attorney",
    "LEWIS SHERIFF": "Sheriff",
    "LEWIS TREASURER": "Treasurer",
    "LEWIS COUNTY COMMISSIONER, DISTRICT 3": "County Commissioner, District 3",
    "DISTRICT COURT DISTRICT COURT JUDGE, DEPT 1": "District Court Judge Dept. 1",
    "DISTRICT COURT DISTRICT COURT JUDGE, DEPT 2": "District Court Judge Dept. 2",
    "PUD DISTRICT-AT-LARGE COMMISSIONER DISTRICT 3": "PUD Commissioner District 3",
    "MORTON CITY PROPOSITION 1 - WITHDRAWAL FROM INTERCOUNTY LIBRARY DISTRICT":
        "Morton Proposition 1 - Withdrawal from Intercounty Library District",
    "FIRE DISTRICT 10 - PACKWOOD PROPOSITION NO. 1 - EMERGENCY MEDICAL SERVICES PROPERTY TAX LEVY":
        "Fire District 10 (Packwood) Proposition 1 - EMS Levy",
}
PARTIES = {"DEMOCRAT": "DEM", "REPUBLICAN": "REP"}


def norm(s):
    return " ".join(str(s).replace("(*)", "").upper().split())


def slug(text):
    return "g22-" + re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:48].rstrip("-")


def nice(name):
    """MEDSL names are upper case: "JONATHAN L MEYER" -> "Jonathan L Meyer"."""
    if name.upper() == "WRITE-IN":
        return "Write-in"
    if name.upper() == "RW BUZZARD":
        return "R.W. Buzzard"
    words = []
    for w in name.title().split():
        if w.startswith("Mc") and len(w) > 2:
            w = "Mc" + w[2:].capitalize()
        words.append(w)
    return " ".join(words)


def round_coords(c):
    return [round_coords(x) for x in c] if isinstance(c[0], (list, tuple)) else [round(c[0], 5), round(c[1], 5)]


def write_boundaries(names):
    gj = json.loads((SITE / "20260804.geojson").read_text(encoding="utf-8"))
    parts = {old: new for new, olds in MERGED.items() for old in olds}
    features, merging = [], defaultdict(list)
    for f in gj["features"]:
        name = f["properties"]["name"]
        if name in parts:
            merging[parts[name]].append(f)
        else:
            features.append(f)
    for name, group in merging.items():
        geom = mapping(unary_union([shape(f["geometry"]) for f in group]))
        props = dict(group[0]["properties"], name=name)
        features.append({"type": "Feature", "properties": props,
                         "geometry": {"type": geom["type"], "coordinates": round_coords(geom["coordinates"])}})
    features.sort(key=lambda f: f["properties"]["num"])
    missing = sorted(set(names) - {f["properties"]["name"] for f in features})
    if missing:
        print(f"no boundary for: {missing}")
    (SITE / "20221108.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": features},
                                                      separators=(",", ":")), encoding="utf-8")


def main(xlsx, medsl):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    names = {r[9]: norm(r[5]) for r in wb["GIS_Precincts_2022Gen"].iter_rows(min_row=2, values_only=True)
             if r[9] and r[9].startswith("LE")}

    # Turnout per Lewis precinct.
    rows = wb["FINAL"].iter_rows(values_only=True)
    header = next(rows)
    i_reg, i_bal = header.index("G22TREGVOT"), header.index("G22TBALCST")
    turnout = {names[r[0]]: (r[i_reg] or 0, r[i_bal] or 0) for r in rows if r[0] and r[0].startswith("LE")}

    # Party preferences, by candidate name, from MEDSL.
    party = {}
    local = []
    with open(medsl, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["county_name"] != "LEWIS" or r["mode"] != "TOTAL":
                continue
            party[norm(r["candidate"]).replace(".", "")] = PARTIES.get(r["party_simplified"], "")
            if r["office"] in LOCAL_CONTESTS:
                local.append(r)

    # State file contests: Lewis precincts, and totals across every county with the contest.
    out = []
    wide = defaultdict(lambda: defaultdict(int))
    counties = defaultdict(set)
    for r in wb["Results"].iter_rows(min_row=2, values_only=True):
        county, prec, code, seq, order, votes, area_type, area, race, choice = r
        key = (area_type, race)
        if key not in STATE_CONTESTS:
            continue
        cid = slug(f"{seq} {area}")
        wide[cid][choice] += votes or 0
        counties[cid].add(county)
        if county == "LE":
            num = str(area).split()[-1]
            out.append({"contest_id": cid, "contest": STATE_CONTESTS[key].format(n=num),
                        "type": "BallotMeasure" if area_type == "SW" and "Advisory" in race else "Candidate",
                        "order": order, "choice": "Write-in" if choice == "WRITE-IN" else choice,
                        "party": party.get(norm(choice).replace(".", ""), ""),
                        "precinct": names[prec], "votes": votes or 0})
    lewis_ids = {r["contest_id"] for r in out}

    for r in local:
        title = LOCAL_CONTESTS[r["office"]]
        measure = r["candidate"] in ("YES", "NO")
        choice = r["candidate"].title() if measure else nice(r["candidate"])
        out.append({"contest_id": slug(title),
                    "contest": title, "type": "BallotMeasure" if measure else "Candidate",
                    # MEDSL has no ballot order: write-ins last, measures Yes first.
                    "order": 97 if choice == "Write-in" else (1 if choice == "Yes" else 2) if measure else 1,
                    "choice": choice, "party": PARTIES.get(r["party_simplified"], ""),
                    "precinct": norm(r["precinct"]), "votes": int(float(r["votes"] or 0))})

    OUT.mkdir(parents=True, exist_ok=True)
    with open(OUT / "turnout.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["precinct", "registered", "ballots_cast"])
        for name in sorted(turnout):
            w.writerow([name, *turnout[name]])
    with open(OUT / "results.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    with open(OUT / "wide.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["contest_id", "counties", "choice", "votes"])
        for cid in sorted(lewis_ids):
            for choice, votes in wide[cid].items():
                w.writerow([cid, len(counties[cid]), "Write-in" if choice == "WRITE-IN" else choice, votes])
    write_boundaries([n for n, (reg, _) in turnout.items() if reg])
    print(f"{len(turnout)} precincts, {len(out)} result rows, {len(lewis_ids)} state-file contests")


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:3]))
